import asyncio
import io
import json
import re

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import config
from .agents.application_writer import make_draft
from .agents.source_critic import check_draft
from .models import DraftRequest, ExportRequest, Profile
from .orchestrator import RUNS, Run, execute
from .tools import llm, parsing

app = FastAPI(title="ScholarSwarm API", version="1.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


@app.get("/api/health")
def health():
    return {"ok": True, "llm": llm.available(), "model": config.LLM_MODEL if llm.available() else None,
            "live_fetch": config.LIVE_FETCH, "search": "tavily" if config.TAVILY_API_KEY else "duckduckgo"}


@app.get("/api/stats")
def stats():
    reg = json.loads((config.DATA_DIR / "opportunities.json").read_text(encoding="utf-8"))["opportunities"]
    return {"registry_sources": len(reg), "countries": len({o["country"] for o in reg}), "agents": 6}




@app.post("/api/upload")
async def upload(kind: str = Form("resume"), file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > 8 * 1024 * 1024:
        raise HTTPException(413, "File too large (8 MB max)")
    text = parsing.extract_text(file.filename, data)
    cg = parsing.detect_cgpa(text)
    return {"kind": kind, "filename": file.filename, "size": len(data), "chars": len(text), "text": text[:60000],
            "skills": parsing.extract_skills(text), "projects": parsing.extract_projects(text),
            "cgpa": cg[0] if cg else None, "cgpa_scale": cg[1] if cg else None}


@app.post("/api/swarm/start")
async def start(profile: Profile):
    run = Run(profile)
    RUNS[run.id] = run
    asyncio.create_task(execute(run))
    return {"run_id": run.id}


@app.get("/api/swarm/{run_id}/stream")
async def stream(run_id: str):
    run = RUNS.get(run_id)
    if not run:
        raise HTTPException(404, "unknown run")

    async def gen():
        i = 0
        while True:
            while i < len(run.events):
                yield f"data: {json.dumps(run.events[i], default=str)}\n\n"
                i += 1
            if run.finished and i >= len(run.events):
                break
            await asyncio.sleep(0.25)

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/swarm/{run_id}/state")
def state(run_id: str):
    run = RUNS.get(run_id)
    if not run:
        raise HTTPException(404, "unknown run")
    return {"agents": list(run.agents.values()), "finished": run.finished, "events": run.events}


@app.get("/api/swarm/{run_id}/result")
def result(run_id: str):
    run = RUNS.get(run_id)
    if not run:
        raise HTTPException(404, "unknown run")
    if not run.finished:
        raise HTTPException(202, "still running")
    return run.result


@app.post("/api/draft")
async def draft(req: DraftRequest):
    run = RUNS.get(req.run_id)
    if not run or not run.finished:
        raise HTTPException(404, "run not found or unfinished")
    target = next((o for o in run.ranked if o["id"] == req.opportunity_id), None)
    if not target:
        raise HTTPException(404, "unknown opportunity")
    p = run.enriched
    text = None
    if llm.available():
        from .agents.application_writer import _facts

        try:
            text = await llm.chat(
                "You draft scholarship Statements of Purpose. STRICT RULES: use ONLY the supplied facts. Never invent publications, awards, "
                "grades, employers or numbers. Wrap project/repo names in backticks. Use [ACTION REQUIRED: ...] where info is missing. "
                "Markdown: '# ' title, Applicant line, 5-6 paragraphs, 450-700 words.",
                f"APPLICANT FACTS:\n{_facts(p)}\n\nPROGRAMME: {target['name']} ({target['provider']}, {target['country']})\n"
                f"Requirements on official page: {target.get('page_requirements')}", 2200, 0.4)
        except Exception:  # noqa: BLE001
            text = None
    d = make_draft(target, p, text)
    d["draft_checks"] = check_draft(d, p)
    run.draft = d
    return d


@app.post("/api/export")
def export(req: ExportRequest):
    clean = re.sub(r"`", "", req.text)
    if req.fmt == "docx":
        import docx

        doc = docx.Document()
        for line in clean.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.strip():
                doc.add_paragraph(line)
        buf = io.BytesIO()
        doc.save(buf)
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": 'attachment; filename="statement_of_purpose.docx"'})
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    story = []
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    for line in clean.splitlines():
        if line.startswith("# "):
            story.append(Paragraph(esc(line[2:]), styles["Heading1"]))
        elif line.strip():
            story.append(Paragraph(esc(line), styles["BodyText"]))
            story.append(Spacer(1, 6))
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": 'attachment; filename="statement_of_purpose.pdf"'})


# Serve the frontend from the same origin (so `uvicorn` alone is enough for the demo)
if config.FRONTEND_DIR.exists():
    app.mount("/", StaticFiles(directory=str(config.FRONTEND_DIR), html=True), name="frontend")
